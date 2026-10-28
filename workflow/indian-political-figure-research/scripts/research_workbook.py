#!/usr/bin/env python3
"""Build a research workbook from durable records; never infer evidence verification.

Requires openpyxl. Input format and review semantics: references/pipeline.md.
"""
import argparse
from collections import Counter, defaultdict
from copy import deepcopy
from datetime import date
import hashlib
import json
import math
from pathlib import Path
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.properties import CalcProperties

FIRST = 6
STATUSES = {'Not reviewed', 'Supported', 'Partially supported', 'Disputed', 'Unsupported'}
NAVY, TEAL, PALE, AMBER = '17354A', '087E8B', 'EAF4F5', 'FFF0D2'


def fingerprint(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True).encode()).hexdigest()


def source_ids(value):
    if not value:
        return []
    parts = [x.strip() for x in value.split(';')]
    if any(not re.fullmatch(r'S\d+', x) for x in parts):
        raise ValueError('Use explicit source IDs, not blanket ranges: ' + value)
    return list(dict.fromkeys(parts))


def index_unique(items, key):
    result = {}
    for item in items:
        ident = item[key]
        if ident in result:
            raise ValueError('Duplicate ID: ' + ident)
        result[ident] = item
    return result


def validate_data(data, retrievals):
    """Validate provenance mechanics, not the truth of the historical narrative."""
    if data['schema_version'] != 1:
        raise ValueError('Unsupported schema version')
    sources = index_unique(data['sources'], 'Source ID')
    claims = index_unique(data['claims'], 'id')
    evidence = index_unique(data['evidence'], 'id')
    fetched = index_unique(retrievals, 'id')
    records = {}
    for name, table in data['tables'].items():
        if len(name) > 31 or any(c in name for c in '[]:*?/\\'):
            raise ValueError('Invalid worksheet name: ' + name)
        for row in table['rows']:
            rid = next(iter(row.values()))
            if rid in records:
                raise ValueError('Duplicate record ID: ' + rid)
            records[rid] = (name, row)
            ids = source_ids(row.get('Source IDs'))
            if not ids and row.get('Claim type') != 'Method / scope':
                raise ValueError('Missing citations: ' + rid)
            if set(ids) - sources.keys():
                raise ValueError('Unknown source on ' + rid)
    for claim in claims.values():
        name, row = records[claim['record_id']]
        if name != claim['sheet'] or row[claim['field']] != claim['statement']:
            raise ValueError('Stale claim statement: ' + claim['id'])
        if fingerprint(claim['statement']) != claim['fingerprint']:
            raise ValueError('Stale evidence fingerprint: ' + claim['id'])
        if claim['status'] not in STATUSES:
            raise ValueError('Invalid review status: ' + claim['id'])
        if set(claim['candidate_source_ids']) - sources.keys():
            raise ValueError('Unknown candidate source: ' + claim['id'])
        supports = []
        for eid in claim['evidence_ids']:
            ev = evidence[eid]
            if ev['claim_id'] != claim['id']:
                raise ValueError('Evidence assigned to wrong claim: ' + eid)
            if ev['relation'] == 'Supports':
                supports.append(ev)
        if claim['status'] == 'Supported':
            if not supports or claim['contradictions']:
                raise ValueError('Unsupported verification or unresolved contradiction: ' + claim['id'])
            # Mirrors and syndicated copies sharing an origin are one source.
            if claim.get('requires_independent_corroboration'):
                origins = {fetched[e['retrieval_id']].get('origin_id') for e in supports}
                if None in origins or len(origins) < 2:
                    raise ValueError('Independent corroboration missing: ' + claim['id'])
    for ev in evidence.values():
        ret = fetched[ev['retrieval_id']]
        if ev['id'] not in claims[ev['claim_id']]['evidence_ids']:
            raise ValueError('Orphan evidence: ' + ev['id'])
        if ev['source_id'] not in sources or ret['source_id'] != ev['source_id']:
            raise ValueError('Evidence/retrieval source mismatch: ' + ev['id'])
        if ret['access_status'] != 'Passage inspected':
            raise ValueError('An unread source cannot support a reviewed claim: ' + ev['id'])
        if not all(ev.get(k) for k in ['locator', 'passage_paraphrase', 'reviewed_by', 'reviewed_at']):
            raise ValueError('Incomplete passage evidence: ' + ev['id'])
    context = data['tables'].get('Context & Constraints', {}).get('rows', [])
    if context:
        ctxids = {r['Context ID'] for r in context}
        for _, row in records.values():
            if set(x.strip() for x in row.get('Context IDs', '').split(';') if x.strip()) - ctxids:
                raise ValueError('Unknown context ID')
        for dimension, expected in data['coverage'].items():
            phases = data['phases'] if dimension == 'Era' else [None]
            for phase in phases:
                rows = [r for r in context if r['Coverage dimension'] == dimension and
                        (phase is None or r['Career phase'] == phase)]
                actual = [r['Coverage item'] for r in rows]
                if len(actual) != len(set(actual)) or set(actual) != set(expected):
                    raise ValueError('Incomplete or duplicate context coverage: ' + str(phase))
    claimed = {c['record_id'] for c in claims.values()}
    if any(rid not in claimed and r.get('Claim type') != 'Method / scope'
           for rid, (_, r) in records.items()):
        raise ValueError('Substantive record missing from claim register')
    return records


def row_status(claims):
    if not claims:
        return 'Method / scope'
    counts = Counter(c['status'] for c in claims)
    if counts['Supported'] == len(claims):
        return 'Supported'
    if counts['Disputed']:
        return 'Disputed'
    if counts['Supported']:
        return f"{counts['Supported']}/{len(claims)} fields supported"
    return 'Not reviewed' if counts['Not reviewed'] == len(claims) else 'Review incomplete'


def make_tables(data, retrievals):
    tables = deepcopy(data['tables'])
    byrecord = defaultdict(list)
    for c in data['claims']:
        byrecord[c['record_id']].append(c)
    citations, used = [], defaultdict(set)
    smap = {s['Source ID']: s for s in data['sources']}
    for name, table in tables.items():
        for row in table['rows']:
            rid = next(iter(row.values()))
            row['Evidence review'] = row_status(byrecord[rid])
            row['All citations'] = rid if row.get('Source IDs') else None
            row['Claim details'] = rid if byrecord[rid] else None
            for sid in source_ids(row.get('Source IDs')):
                used[sid].add(name)
                citations.append({'Citation ID': f'{rid}:{sid}', 'Record ID': rid,
                                  'Source ID': sid, 'Title': smap[sid]['Title'],
                                  'URL': smap[sid]['URL'], 'Access status': smap[sid]['Access status'],
                                  'Sheet': name, 'Meaning': 'Candidate reference; inspect Claims and Evidence for checked support.'})
        table['lead_columns'] += ['Evidence review', 'All citations', 'Claim details']
    sources = deepcopy(data['sources'])
    for s in sources:
        s['Used in'] = '; '.join(n for n in tables if n in used[s['Source ID']])
    tables['Sources'] = {'lead_columns': ['Source ID', 'Title', 'Access status', 'URL'], 'rows': sources}
    tables['Claims'] = {'lead_columns': ['Claim ID', 'Record ID', 'Statement', 'Review status', 'Evidence'], 'rows': [
        {'Claim ID': c['id'], 'Record ID': c['record_id'], 'Statement': c['statement'],
         'Review status': c['status'], 'Evidence': '; '.join(c['evidence_ids']) or None,
         'Priority': c['priority'], 'Field': c['field'], 'Claim type': c['claim_type'],
         'Candidate sources': '; '.join(c['candidate_source_ids']), 'All citations': c['record_id'],
         'Contradictions': '; '.join(c['contradictions']), 'Review note': c['review_note'],
         'Sheet': c['sheet']} for c in data['claims']]}
    tables['Evidence'] = {'lead_columns': ['Evidence ID', 'Claim ID', 'Source ID', 'Passage paraphrase', 'URL'], 'rows': [
        {'Evidence ID': e['id'], 'Claim ID': e['claim_id'], 'Source ID': e['source_id'],
         'Passage paraphrase': e['passage_paraphrase'], 'URL': smap[e['source_id']]['URL'],
         'Locator': e['locator'], 'Relation': e['relation'], 'Retrieval ID': e['retrieval_id'],
         'Reviewer': e['reviewed_by'], 'Reviewed at': e['reviewed_at'], 'Note': e['note']}
        for e in data['evidence']]}
    tables['Citations'] = {'lead_columns': ['Citation ID', 'Record ID', 'Source ID', 'Title', 'URL'], 'rows': citations}
    tables['Research Queue'] = {'lead_columns': ['Gap ID', 'Priority', 'Topic', 'Gap', 'Next action'], 'rows': [
        {'Gap ID': g['id'], 'Priority': g['priority'], 'Topic': g['topic'], 'Gap': g['gap'],
         'Next action': g['next_action'], 'Status': g['status'], 'Records': '; '.join(g['record_ids'])}
        for g in data['gaps']]}
    counts = Counter(c['status'] for c in data['claims'])
    els = tables.get('Elections', {}).get('rows', [])
    direct = [r for r in els if r['House / route'].endswith((' / direct', ' / by-election'))]
    metrics = [('Supported field claims', counts['Supported']), ('Registered field claims', len(data['claims'])),
               ('High-priority claims awaiting support', sum(c['priority'] == 'High' and c['status'] != 'Supported' for c in data['claims'])),
               ('Direct contests missing candidate votes', sum(r['Candidate votes'] is None for r in direct)),
               ('Direct contests missing vote share', sum(r['Vote share'] is None for r in direct))]
    overview = tables['Overview']
    idcol = next(iter(overview['rows'][0]))
    for i, (topic, value) in enumerate(metrics, 1):
        overview['rows'].append({idcol: f'QA{i:02d}', 'Topic': topic, 'Finding / value': value,
                                'Section': 'Evidence coverage', 'Claim type': 'Derived count',
                                'Evidence review': 'Workbook calculation'})
    for ident, topic, message in [
        ('GUIDE', 'Context guidance', data['guidance']),
        ('REVIEW', 'How to read evidence coverage', 'The register tracks field-level assertions, some compound. Supported requires checking the entire field. Migrated assertions remain Not reviewed until passage evidence is recorded; this is not a finding that they are false.'),
        ('QA', 'Validation scope', 'Structural checks, arithmetic, evidence review, application recalculation and visual inspection are separate. See validation.json for actual results. Formulas recalculate on opening; no synthetic formula caches are inserted.')]:
        overview['rows'].append({idcol: ident, 'Topic': topic, 'Finding / value': message,
                                'Claim type': 'Method / scope', 'Evidence review': 'Method / scope'})
    return tables, metrics


def columns_for(table):
    columns = list(table['lead_columns'])
    for row in table['rows']:
        columns.extend(k for k in row if k not in columns)
    return columns


def cell_width(column):
    if column.endswith(' ID') or column in ['All citations', 'Claim details', 'Evidence']:
        return 16
    if column in ['Year', 'Result', 'Priority']:
        return 13
    if column in ['Review status', 'Evidence review', 'Access status', 'Claim type']:
        return 24
    if column in ['Finding / value', 'Finding', 'Statement', 'Passage paraphrase', 'Gap', 'Next action']:
        return 58
    if 'date' in column.lower() or column in ['Candidate votes', 'Opponent votes', 'Vote share', 'Signed vote margin']:
        return 18
    return 34 if column in ['URL', 'Topic', 'Title', 'Event', 'Theme', 'Position / institution', 'Initiative / decision'] else 44


def build(data, retrievals, output, report_path):
    records = validate_data(data, retrievals)
    tables, metrics = make_tables(data, retrievals)
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title = data['subject'] + ' — Political Research'
    wb.properties.creator = 'Political Research by ChatGPT'
    wb.calculation = CalcProperties(fullCalcOnLoad=True, forceFullCalc=True)
    loc, cols, formula_checks = {}, {}, []
    for name, table in tables.items():
        cols[name] = columns_for(table)
        for ri, row in enumerate(table['rows'], FIRST):
            loc[(name, str(row[cols[name][0]]))] = ri
    record_loc = {rid: (name, loc[(name, rid)]) for rid, (name, _) in records.items()}
    first_claim, first_citation = {}, {}
    for row in tables['Claims']['rows']:
        first_claim.setdefault(row['Record ID'], loc[('Claims', row['Claim ID'])])
    for row in tables['Citations']['rows']:
        first_citation.setdefault(row['Record ID'], loc[('Citations', row['Citation ID'])])

    def link(cell, sheet, row):
        cell.hyperlink = f"#'{sheet.replace(chr(39), chr(39)*2)}'!A{row}"
        cell.font = Font(name='Arial', size=11, color=TEAL, underline='single')

    for sheetno, (name, table) in enumerate(tables.items(), 1):
        ws = wb.create_sheet(name)
        cs, lead = cols[name], len(table['lead_columns'])
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 85
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=lead)
        ws.cell(1, 1, data['subject'] + ' | ' + name)
        ws['A1'].font = Font(name='Arial', size=18, bold=True, color='FFFFFF')
        ws['A1'].fill = PatternFill('solid', fgColor=NAVY)
        ws.row_dimensions[1].height = 36
        for r, text in [(2, 'As of ' + data['as_of'] + ' | Expand the + column group for supporting detail.'),
                        (3, 'Citation IDs are references, not proof. Read evidence status and claim qualifications.')]:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=lead)
            ws.cell(r, 1, text).font = Font(name='Arial', size=11, color='647786')
            ws.row_dimensions[r].height = 20
        for ci, c in enumerate(cs, 1):
            cell = ws.cell(5, ci, c)
            cell.font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor=TEAL)
            cell.alignment = Alignment(wrap_text=True, vertical='center')
            ws.column_dimensions[letter(ci)].width = cell_width(c)
        ws.row_dimensions[5].height = 36
        if len(cs) > lead:
            for ci in range(lead+1, len(cs)+1):
                ws.column_dimensions[letter(ci)].hidden = True
                ws.column_dimensions[letter(ci)].outlineLevel = 1
            ws.sheet_properties.outlinePr.summaryRight = False
            ws.column_dimensions[letter(lead)].collapsed = True
        for ri, row in enumerate(table['rows'], FIRST):
            lines = 1
            for ci, c in enumerate(cs, 1):
                value = row.get(c)
                if c in table.get('date_columns', []) and value:
                    value = date.fromisoformat(value)
                cell = ws.cell(ri, ci, value)
                # Research text must never become an executable spreadsheet formula.
                if isinstance(value, str) and value.startswith('='):
                    cell.data_type = 's'
                cell.font = Font(name='Arial', size=11, color='203746')
                cell.fill = PatternFill('solid', fgColor=PALE if ri % 2 == 0 else 'FFFFFF')
                cell.alignment = Alignment(wrap_text=True, vertical='top')
                if isinstance(value, date):
                    cell.number_format = 'd mmm yyyy'
                elif c == 'Vote share':
                    cell.number_format = '0.00%'
                elif isinstance(value, (int, float)) and c != 'Year':
                    cell.number_format = '#,##0'
                if c in ['Evidence review', 'Review status'] and value not in ['Supported', 'Method / scope', 'Workbook calculation']:
                    cell.fill = PatternFill('solid', fgColor=AMBER)
                if c == 'URL' and value:
                    cell.hyperlink = value
                    cell.font = Font(name='Arial', size=11, color=TEAL, underline='single')
                elif c == 'All citations' and value in first_citation:
                    link(cell, 'Citations', first_citation[value])
                    cell.value = 'All citations'
                elif c == 'Claim details' and value in first_claim:
                    link(cell, 'Claims', first_claim[value])
                    cell.value = 'Claims'
                elif c == 'Record ID' and value in record_loc:
                    link(cell, *record_loc[value])
                elif c == 'Source ID' and name != 'Sources' and value:
                    link(cell, 'Sources', loc[('Sources', value)])
                elif c == 'Claim ID' and name != 'Claims' and value:
                    link(cell, 'Claims', loc[('Claims', value)])
                elif c == 'Evidence' and value:
                    link(cell, 'Evidence', loc[('Evidence', value.split(';')[0])])
                if value is not None:
                    lines = max(lines, sum(max(1, math.ceil(len(x) / (cell_width(c)*0.85))) for x in str(value).splitlines()))
            # This estimate sets workbook dimensions; it is not a rendering check.
            ws.row_dimensions[ri].height = min(405, max(44, lines*14+10))
        if table['rows']:
            ref = f'A5:{letter(len(cs))}{ws.max_row}'
            tab = Table(displayName=f'Research_{sheetno}', ref=ref)
            tab.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
            ws.add_table(tab)
            ws.auto_filter.ref = ref
        ws.freeze_panes = 'C6'
        ws.print_title_rows = '1:5'
        ws.print_area = f'A1:{letter(lead)}{ws.max_row}'
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0

    def address(sheet, column, row=None):
        col = letter(cols[sheet].index(column)+1)
        return col + str(row) if row else col

    def add_formula(sheet, row, column, formula, expected):
        cell = address(sheet, column, row)
        wb[sheet][cell] = formula
        formula_checks.append({'sheet': sheet, 'cell': cell, 'formula': formula, 'expected': expected,
                               'record_id': str(wb[sheet].cell(row, 1).value), 'column': column})

    els = tables.get('Elections', {}).get('rows', [])
    for ri, row in enumerate(els, FIRST):
        if isinstance(row['Candidate votes'], int) and isinstance(row['Opponent votes'], int):
            add_formula('Elections', ri, 'Signed vote margin',
                        '=' + address('Elections', 'Candidate votes', ri) + '-' + address('Elections', 'Opponent votes', ri),
                        row['Candidate votes']-row['Opponent votes'])
    lastclaim = FIRST+len(data['claims'])-1
    status = address('Claims', 'Review status')
    countcol = address('Claims', 'Claim ID')
    add_formula('Overview', loc[('Overview', 'QA01')], 'Finding / value',
                f'=COUNTIF(Claims!{status}{FIRST}:{status}{lastclaim},"Supported")', metrics[0][1])
    add_formula('Overview', loc[('Overview', 'QA02')], 'Finding / value',
                f'=COUNTA(Claims!{countcol}{FIRST}:{countcol}{lastclaim})', metrics[1][1])
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    check = load_workbook(output)
    if check.sheetnames != list(tables):
        raise ValueError('Saved workbook lost worksheets')
    link_count = 0
    for ws in check:
        if ws.freeze_panes != 'C6' or (ws.max_row >= FIRST and len(ws.tables) != 1):
            raise ValueError('Navigation/table validation failed: ' + ws.title)
        for row in ws:
            for cell in row:
                if cell.data_type == 'e':
                    raise ValueError('Spreadsheet error: ' + ws.title + cell.coordinate)
                if cell.hyperlink:
                    link_count += 1
                    target = cell.hyperlink.target
                    if target.startswith('#'):
                        match = re.fullmatch(r"#'((?:[^']|'')+)'!([A-Z]+\d+)", target)
                        if not match or check[match[1].replace("''", "'")][match[2]].value is None:
                            raise ValueError('Broken internal hyperlink: ' + target)
    report = {'subject': data['subject'], 'file': str(output),
              'workbook_sha256': hashlib.sha256(output.read_bytes()).hexdigest(),
              'checks': {'structure': 'PASS', 'provenance_integrity': 'PASS',
                         'arithmetic': 'PASS: expectations calculated independently in Python',
                         'application_recalculation': 'NOT RUN', 'visual_inspection': 'NOT RUN'},
              'evidence': {'registered_fields': len(data['claims']), 'status_counts': dict(Counter(c['status'] for c in data['claims'])),
                           'note': 'These are review units, not a percentage of historical truth or complete atomic-claim coverage.'},
              'sheets': {n: len(t['rows']) for n, t in tables.items()},
              'hyperlinks_checked': link_count, 'formula_checks': formula_checks,
              'limitations': ['No formula caches manufactured.', 'Row-height estimates are not a visual check.',
                              'Source presence does not certify factual support; inherited claim types await review.']}
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(report, indent=2)+'\n')
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('data', type=Path)
    parser.add_argument('--output', required=True, type=Path)
    parser.add_argument('--report', type=Path)
    args = parser.parse_args()
    data = json.loads(args.data.read_text())
    retrievals = json.loads(args.data.with_name('retrievals.json').read_text())
    result = build(data, retrievals, args.output, args.report or args.data.with_name('validation.json'))
    print(json.dumps({k: result[k] for k in ['checks', 'evidence', 'sheets']}, indent=2))


if __name__ == '__main__':
    main()
