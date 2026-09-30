#!/usr/bin/env python3
"""Compare actual spreadsheet-engine results with independent expected values."""
import argparse
import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook


def compare(workbook, checks):
    values = load_workbook(workbook, data_only=True)
    formulas = load_workbook(workbook, data_only=False)
    results = []
    for check in checks:
        matches = []
        for ws in values:
            headers = [(c.row, c.column) for row in ws.iter_rows(max_row=min(20, ws.max_row))
                       for c in row if c.value == check['column']]
            if not headers:
                continue
            for hr, col in headers:
                for row in ws.iter_rows(min_row=hr+1):
                    if any(str(c.value) == check['record_id'] for c in row):
                        target = ws.cell(row[0].row, col)
                        formula = formulas[ws.title][target.coordinate]
                        if formula.data_type == 'f':
                            matches.append((ws.title, target.coordinate, target.value))
        if len(matches) != 1:
            raise ValueError(f"Expected one engine formula for {check['record_id']}/{check['column']}; found {len(matches)}")
        sheet, address, value = matches[0]
        if not isinstance(value, (int, float)) or abs(value-check['expected']) > 1e-8:
            raise ValueError(f"Engine result mismatch for {check['record_id']}: {value!r} != {check['expected']}")
        results.append({'record_id': check['record_id'], 'sheet': sheet, 'cell': address, 'engine_value': value})
    return results


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('report', type=Path)
    p.add_argument('original', type=Path)
    p.add_argument('recalculated', type=Path)
    p.add_argument('--engine', required=True)
    args = p.parse_args()
    report = json.loads(args.report.read_text())
    if hashlib.sha256(args.original.read_bytes()).hexdigest() != report['workbook_sha256']:
        raise ValueError('Original workbook changed after validation; rebuild and rerun checks')
    try:
        results = compare(args.recalculated, report['formula_checks'])
        report['application'] = {'engine': args.engine, 'export_sha256': hashlib.sha256(args.recalculated.read_bytes()).hexdigest(), 'results': results}
        report['checks']['application_recalculation'] = 'PASS'
    except Exception as exc:
        report['checks']['application_recalculation'] = 'FAIL'
        report['application'] = {'engine': args.engine, 'error': str(exc)}
        raise
    finally:
        args.report.write_text(json.dumps(report, indent=2)+'\n')
    print(f'PASS: {len(results)} formula results checked in {args.engine}')


if __name__ == '__main__':
    main()
